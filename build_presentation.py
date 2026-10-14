import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

SRC_PPT = '/Users/chaitanya/Desktop/Waste2Worth_AI/Jugaad_Engineering_Contest_Blueprint.pptx'
OUT_PPT = '/Users/chaitanya/Desktop/Waste2Worth_AI/Waste2Worth_Sortify_Presentation.pptx'
RESULTS_DIR = '/Users/chaitanya/Desktop/Waste2Worth_AI/results'

prs = Presentation(SRC_PPT)

# Delete Slide 8 (General Notes / Guidelines)
if len(prs.slides) > 7:
    rId = prs.slides._sldIdLst[7].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[7]

# Brand & Blueprint Theme Colors
COLOR_AMBER = RGBColor(0xF0, 0xA2, 0x02)       # Accent gold/amber
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)       # Pure white
COLOR_ICE_BLUE = RGBColor(0xDC, 0xEE, 0xF5)    # Technical light blue
COLOR_CYAN = RGBColor(0x9F, 0xC4, 0xD6)        # Muted cyan/blueprint blue
COLOR_HIGHLIGHT = RGBColor(0x00, 0xE5, 0xFF)   # High-visibility neon cyan
COLOR_BOX_BG = RGBColor(0x07, 0x1C, 0x2C)      # Deep blueprint container navy
COLOR_CARD_BORDER = RGBColor(0xF0, 0xA2, 0x02) # Amber border
COLOR_FRAME_BORDER = RGBColor(0x24, 0x48, 0x66)# Subtle steel blue border

def format_item_text(shape, bold_title, desc_text, font_size=11.5, width=None):
    if width is not None:
        shape.width = width
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    
    p = tf.paragraphs[0]
    p.text = ''
    p.line_spacing = 1.05
    
    r1 = p.add_run()
    r1.text = bold_title + " — "
    r1.font.name = 'Calibri'
    r1.font.size = Pt(font_size)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_AMBER
    
    r2 = p.add_run()
    r2.text = desc_text
    r2.font.name = 'Calibri'
    r2.font.size = Pt(font_size)
    r2.font.bold = False
    r2.font.color.rgb = COLOR_ICE_BLUE

# ==========================================
# SLIDE 1: TITLE (Sheet 01 / 07)
# ==========================================
s1 = prs.slides[0]
# Shape 27 is TEAM ID text box ("—"), Shape 29 is PROJECT TITLE ("—")
for s in s1.shapes:
    if s.has_text_frame and s.text_frame.text.strip() == "—":
        if s.left < 5000000: # Left quadrant: TEAM ID
            s.top = 5540000
            s.height = 320000
            tf = s.text_frame
            tf.word_wrap = True
            tf.margin_top = Inches(0.01)
            tf.margin_bottom = Inches(0.01)
            p = tf.paragraphs[0]
            p.text = "SORTIFY"
            p.font.name = 'Calibri'
            p.font.size = Pt(19)
            p.font.bold = True
            p.font.color.rgb = COLOR_WHITE
        else: # Right quadrant: PROJECT TITLE
            s.top = 5540000
            s.height = 320000
            tf = s.text_frame
            tf.word_wrap = True
            tf.margin_top = Inches(0.01)
            tf.margin_bottom = Inches(0.01)
            p = tf.paragraphs[0]
            p.text = "WASTE2WORTH AI"
            p.font.name = 'Calibri'
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = COLOR_WHITE

# ==========================================
# SLIDE 2: PROBLEM STATEMENT (Sheet 02 / 07)
# ==========================================
s2 = prs.slides[1]
slide2_data = [
    ("Unsegregated & Misclassified Waste", "Municipal and campus solid waste is dumped in mixed streams without source segregation, causing over 80% of recyclable material to end up in landfills."),
    ("Where & When It Occurs", "High-density university campuses, public community bins, and municipal streets where waste piles accumulate and overflow before authorities notice."),
    ("Who Is Affected", "Citizens facing unhygienic surroundings; municipal sanitation crews exposed to hazardous mixed waste; recycling centers burdened with contaminated batches."),
    ("Current Situation & Non-Jugaad Methods", "Relies on manual, error-prone human sorting at dump sites or expensive industrial optical sorters ($50k+) that are completely inaccessible for field deployment."),
    ("Why It Is Worth Solving", "Lack of traceability causes multi-day collection backlogs; improper identification allows toxic e-waste and medical waste to contaminate compostable soil."),
    ("Key Need: Connected Civic Workflow", "Critical demand for a zero-hardware-cost platform bridging citizens and collectors: instant smartphone scanning, GPS geotagging, verified dispatch, and tracking.")
]

s2_boxes = [s2.shapes[20], s2.shapes[23], s2.shapes[26], s2.shapes[29], s2.shapes[32], s2.shapes[35]]
for box, (title, desc) in zip(s2_boxes, slide2_data):
    format_item_text(box, title, desc, font_size=12.5)

# ==========================================
# SLIDE 3: PROPOSED SOLUTION (Sheet 03 / 07)
# ==========================================
s3 = prs.slides[2]
slide3_data = [
    ("Solution Overview: Waste2Worth AI", "An intelligent, closed-loop civic waste identification and collection platform transforming any smartphone browser into an instant AI waste-sorting scanner."),
    ("Addressing The Problem at Source", "Replaces manual guessing with instant 7-class deep learning classification, provides prediction confidence, and captures high-precision GPS coordinates for dispatch."),
    ("Zero-Hardware \"Jugaad\" Architecture", "100% software-driven resourcefulness — uses existing citizen smartphone cameras, native browser geolocation, and open-source models with zero capital expenditure."),
    ("Core Platform Capabilities", "Synchronized 3-tier civic ecosystem: Citizen Scanner (scan, classify, report), Admin Command Center (triage, assign), and Collector Portal (navigate, proof)."),
    ("End-to-End Operational Lifecycle", "Citizen scans waste → AI identifies category & confidence → GPS report submitted → Admin verifies → Collector assigned → Waste collected → Evidence verified."),
    ("System Architecture & Data Flow", "Citizen/Admin/Collector Responsive UI → FastAPI REST Backend → PyTorch ResNet-18 Inference Service (Apple Silicon MPS) → PostgreSQL Relational Database.")
]

s3_boxes = [s3.shapes[20], s3.shapes[23], s3.shapes[26], s3.shapes[29], s3.shapes[32], s3.shapes[35]]
for box, (title, desc) in zip(s3_boxes, slide3_data):
    format_item_text(box, title, desc, font_size=12.5)

# ==========================================
# SLIDE 4: JUGAAD INNOVATION & UNIQUENESS (Sheet 04 / 07)
# ==========================================
s4 = prs.slides[3]
slide4_data = [
    ("What Makes It Different", "Bridges the gap between advanced deep learning and grassroots civic action — tightly integrates real-time AI classification directly with municipal collection logistics."),
    ("Existing Solutions & Current Alternatives", "Industrial conveyor-belt optical sorters ($50k–$100k), expensive smart bins with embedded load/IR sensors ($1,500–$5,000/bin), and slow municipal phone helplines."),
    ("Limitations of Existing Approaches", "Smart bins are cost-prohibitive, fragile, and prone to vandalization; manual helplines lack image verification, category classification, and accurate geotagging."),
    ("The Core \"Jugaad\" — Hardware-Free AI", "Converts everyday citizen smartphone cameras into intelligent waste sensors without requiring any specialized scanning hardware, external attachments, or costly sensors."),
    ("Resourcefulness & Engineering Ingenuity", "Combines WebRTC browser camera access, native GPS, and an optimized PyTorch ResNet-18 model trained locally with MPS acceleration, lowering deployment cost to near zero.")
]

s4_boxes = [s4.shapes[20], s4.shapes[23], s4.shapes[26], s4.shapes[29], s4.shapes[32]]
for box, (title, desc) in zip(s4_boxes, slide4_data):
    format_item_text(box, title, desc, font_size=13.0)

# ==========================================
# SLIDE 5: MATERIALS & TECHNICAL APPROACH (Sheet 05 / 07)
# ==========================================
s5 = prs.slides[4]
slide5_data = [
    ("Core AI Model: PyTorch ResNet-18", "11.7M parameter residual neural network fine-tuned for 7 waste classes, achieving high accuracy with ultra-low latency (~22ms inference on Apple Silicon MPS)."),
    ("Dataset: 14,279 Curated Images", "14,279 total images (11,423 training / 2,856 validation) across 7 classes: Cardboard, E-waste, Glass, Medical (biohazard priority), Metal, Paper, and Plastic."),
    ("Computer Vision & Preprocessing", "OpenCV & Torchvision pipeline: WebRTC camera capture → 256×256 bilinear resize → 224×224 center crop → ImageNet tensor normalization → Softmax classification."),
    ("Full-Stack Software Architecture", "FastAPI asynchronous Python backend with JWT multi-role auth; responsive HTML5 web frontend; PostgreSQL relational database for users, reports, and evidence."),
    ("Hardware Acceleration & Training", "Trained locally on Apple Silicon MPS (Metal Performance Shaders) with AdamW optimizer, cosine annealing schedule, and automatic CUDA/CPU fallback."),
    ("Sensors & Geolocation Integration", "Browser Camera API (WebRTC video streaming) and HTML5 Geolocation API (high-accuracy GPS latitude/longitude) for automated spatial incident mapping.")
]

split_width = 5700000 # Left column width (EMU)
s5_boxes = [s5.shapes[20], s5.shapes[23], s5.shapes[26], s5.shapes[29], s5.shapes[32], s5.shapes[35]]
for box, (title, desc) in zip(s5_boxes, slide5_data):
    format_item_text(box, title, desc, font_size=11.5, width=split_width)

# Right side of Slide 5: 7 Classes Card & Accuracy Plot
panel5_left = 7100000
panel5_width = 4500000

card1 = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, panel5_left, 2050000, panel5_width, 1400000)
card1.fill.solid()
card1.fill.fore_color.rgb = COLOR_BOX_BG
card1.line.color.rgb = COLOR_CARD_BORDER
card1.line.width = Pt(1)

tf1 = card1.text_frame
tf1.word_wrap = True
tf1.margin_left = Inches(0.12)
tf1.margin_right = Inches(0.12)
tf1.margin_top = Inches(0.08)
tf1.margin_bottom = Inches(0.08)

p1 = tf1.paragraphs[0]
p1.text = "7 TARGET WASTE CLASSES (14,279 IMAGES)"
p1.font.name = 'Courier New'
p1.font.size = Pt(10)
p1.font.bold = True
p1.font.color.rgb = COLOR_AMBER

p1_sub = tf1.add_paragraph()
p1_sub.text = "1. Cardboard  ·  2. E-waste  ·  3. Glass  ·  4. Medical"
p1_sub.font.name = 'Calibri'
p1_sub.font.size = Pt(11)
p1_sub.font.bold = True
p1_sub.font.color.rgb = COLOR_WHITE

p1_sub2 = tf1.add_paragraph()
p1_sub2.text = "5. Metal  ·  6. Paper  ·  7. Plastic"
p1_sub2.font.name = 'Calibri'
p1_sub2.font.size = Pt(11)
p1_sub2.font.bold = True
p1_sub2.font.color.rgb = COLOR_WHITE

p1_desc = tf1.add_paragraph()
p1_desc.text = "Inference: ~22ms / frame on Apple Silicon MPS (Torchvision ResNet-18)"
p1_desc.font.name = 'Courier New'
p1_desc.font.size = Pt(9.5)
p1_desc.font.color.rgb = COLOR_CYAN

# Add Accuracy Curve Image on Slide 5
acc_img_path = os.path.join(RESULTS_DIR, 'accuracy_curve.png')
if os.path.exists(acc_img_path):
    img_shape = s5.shapes.add_picture(acc_img_path, panel5_left, 3550000, width=panel5_width, height=2300000)
    # Caption
    caption = s5.shapes.add_textbox(panel5_left, 5870000, panel5_width, 240000)
    c_p = caption.text_frame.paragraphs[0]
    c_p.text = "Fig 5.1: Actual ResNet-18 Validation Accuracy Curve (10 Epochs, MPS)"
    c_p.font.name = 'Courier New'
    c_p.font.size = Pt(9)
    c_p.font.color.rgb = COLOR_CYAN
    c_p.alignment = PP_ALIGN.CENTER

# ==========================================
# SLIDE 6: RESULTS / FEASIBILITY / IMPACT (Sheet 06 / 07)
# ==========================================
s6 = prs.slides[5]
# Concise 2-line descriptions for tight vertical rhythm
slide6_data = [
    ("Validated AI Performance: 87.89%", "Tested on 2,856 unseen validation images: 87.89% accuracy, 0.8774 Macro F1, 0.8787 Weighted F1. Top precision on Cardboard (94.9%) and Paper (88.3%)."),
    ("True Jugaad Feasibility: ₹0 Hardware Cost", "Zero expenditure on specialized hardware; operates entirely on standard citizen smartphones and open-source backend stack."),
    ("Resource & Material Availability", "14,279 open-source images curated; web application accessible immediately via mobile browser without app store installation or high-end specs."),
    ("Major Challenges & Mitigations", "Mitigated variable lighting and cluttered backgrounds via aggressive data augmentation (flips, rotations, color jitter) and >70% confidence gating."),
    ("Campus & Community Impact", "Empowers citizens to report waste in seconds; reduces municipal response time from days to hours; pre-identifies hazardous medical and e-waste."),
    ("Environmental & Circular Benefit", "Maximizes source segregation; diverts recyclable paper, glass, metal, cardboard, and plastic from landfills while safely isolating e-waste."),
    ("Measurable Project Outcomes", "14,279 images | 87.89% Accuracy | 0.8774 Macro F1 | 0.8787 Weighted F1 | ~22ms latency | 100% geotagged incident traceability.")
]

s6_boxes = [s6.shapes[20], s6.shapes[23], s6.shapes[26], s6.shapes[29], s6.shapes[32], s6.shapes[35], s6.shapes[38]]
for box, (title, desc) in zip(s6_boxes, slide6_data):
    format_item_text(box, title, desc, font_size=11.0, width=split_width)

# Right side of Slide 6: Metrics Card + Normalized Confusion Matrix
panel6_left = 7100000
panel6_width = 4500000

card6 = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, panel6_left, 2050000, panel6_width, 1300000)
card6.fill.solid()
card6.fill.fore_color.rgb = COLOR_BOX_BG
card6.line.color.rgb = COLOR_CARD_BORDER
card6.line.width = Pt(1)

tf6 = card6.text_frame
tf6.word_wrap = True
tf6.margin_left = Inches(0.12)
tf6.margin_right = Inches(0.12)
tf6.margin_top = Inches(0.08)
tf6.margin_bottom = Inches(0.08)

p6 = tf6.paragraphs[0]
p6.text = "OFFICIAL MODEL EVALUATION METRICS"
p6.font.name = 'Courier New'
p6.font.size = Pt(10)
p6.font.bold = True
p6.font.color.rgb = COLOR_AMBER

p6_main = tf6.add_paragraph()
r_acc = p6_main.add_run()
r_acc.text = "87.89% "
r_acc.font.name = 'Calibri'
r_acc.font.size = Pt(22)
r_acc.font.bold = True
r_acc.font.color.rgb = COLOR_HIGHLIGHT

r_lbl = p6_main.add_run()
r_lbl.text = "Validation Accuracy (N=2,856)"
r_lbl.font.name = 'Calibri'
r_lbl.font.size = Pt(12)
r_lbl.font.color.rgb = COLOR_WHITE

p6_stats = tf6.add_paragraph()
p6_stats.text = "Dataset: 14,279 Images (11,423 Train / 2,856 Val)\nMacro F1: 0.8774  |  Weighted F1: 0.8787  |  Classes: 7"
p6_stats.font.name = 'Courier New'
p6_stats.font.size = Pt(9.5)
p6_stats.font.color.rgb = COLOR_CYAN

# Add Confusion Matrix Image on Slide 6
cm_img_path = os.path.join(RESULTS_DIR, 'confusion_matrix_normalized.png')
if os.path.exists(cm_img_path):
    cm_width = 3300000
    cm_height = 2500000
    cm_left = panel6_left + int((panel6_width - cm_width) / 2)
    s6.shapes.add_picture(cm_img_path, cm_left, 3450000, width=cm_width, height=cm_height)
    
    caption6 = s6.shapes.add_textbox(panel6_left, 5970000, panel6_width, 240000)
    c6_p = caption6.text_frame.paragraphs[0]
    c6_p.text = "Fig 6.1: Actual Validation Confusion Matrix (7 Classes, N=2,856)"
    c6_p.font.name = 'Courier New'
    c6_p.font.size = Pt(9)
    c6_p.font.color.rgb = COLOR_CYAN
    c6_p.alignment = PP_ALIGN.CENTER

# ==========================================
# SLIDE 7: PROTOTYPE & FUTURE SCOPE (Sheet 07 / 07)
# ==========================================
s7 = prs.slides[6]
slide7_data = [
    ("Functional Multi-Role Prototype", "Complete operational web platform with 3 dedicated portals: Citizen Scanner (camera scan & GPS), Admin Command (triage & dispatch), and Collector Portal (navigation & proof)."),
    ("Round 2 Live Demonstration Plan", "Live physical demonstration of scanning real scrap objects using mobile camera, sub-second AI inference, instant GPS report submission, and collector task assignment."),
    ("Offline Build & Prototyping Roadmap", "Round 2 development: incorporate real campus waste sample data, containerize deployment with Docker Compose, and configure automated alert triggers."),
    ("Complete Operational Workflow", "Citizen Scans Waste → AI Classifies Category & Confidence → GPS Geotagged Report → Admin Verifies → Collector Dispatched → Evidence Photo Uploaded → Completed."),
    ("Future Technical Enhancements", "Upgrade to YOLOv8 object detection for multi-item recognition in a single frame; expand to 15+ sub-categories; integrate AI-based GIS route optimization for collectors."),
    ("Scalability Across Campuses & Cities", "Lightweight progressive web app (PWA) architecture and low server footprint allow frictionless scaling from campus to city-wide municipal wards with minimal cost."),
    ("Wider Civic & Circular Economy Adoption", "Multilingual audio/visual support for sanitation workers, gamified eco-rewards points to drive citizen participation, and automated municipal diversion analytics.")
]

s7_boxes = [s7.shapes[20], s7.shapes[23], s7.shapes[26], s7.shapes[29], s7.shapes[32], s7.shapes[35], s7.shapes[38]]
for box, (title, desc) in zip(s7_boxes, slide7_data):
    format_item_text(box, title, desc, font_size=11.0, width=split_width)

# Right side of Slide 7: 3 Synchronized Portals Card & End-to-End Diagram Card
panel7_left = 7100000
panel7_width = 4500000

card7_1 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, panel7_left, 2050000, panel7_width, 1850000)
card7_1.fill.solid()
card7_1.fill.fore_color.rgb = COLOR_BOX_BG
card7_1.line.color.rgb = COLOR_CARD_BORDER
card7_1.line.width = Pt(1)

tf7_1 = card7_1.text_frame
tf7_1.word_wrap = True
tf7_1.margin_left = Inches(0.12)
tf7_1.margin_right = Inches(0.12)
tf7_1.margin_top = Inches(0.08)
tf7_1.margin_bottom = Inches(0.08)

p7_1 = tf7_1.paragraphs[0]
p7_1.text = "THE 3 SYNCHRONIZED USER INTERFACES"
p7_1.font.name = 'Courier New'
p7_1.font.size = Pt(10)
p7_1.font.bold = True
p7_1.font.color.rgb = COLOR_AMBER

p7_c = tf7_1.add_paragraph()
r = p7_c.add_run()
r.text = "1. CITIZEN PORTAL: "
r.font.bold = True
r.font.color.rgb = COLOR_WHITE
r.font.size = Pt(10)
r2 = p7_c.add_run()
r2.text = "Mobile WebRTC camera waste scan, real-time AI classification & confidence, auto GPS capture, report submission & tracking."
r2.font.color.rgb = COLOR_ICE_BLUE
r2.font.size = Pt(9.5)

p7_a = tf7_1.add_paragraph()
r = p7_a.add_run()
r.text = "2. ADMIN COMMAND: "
r.font.bold = True
r.font.color.rgb = COLOR_WHITE
r.font.size = Pt(10)
r2 = p7_a.add_run()
r2.text = "Incident triage queue, report verification/audit, collector task assignment, landfill diversion metrics, and waste analytics."
r2.font.color.rgb = COLOR_ICE_BLUE
r2.font.size = Pt(9.5)

p7_w = tf7_1.add_paragraph()
r = p7_w.add_run()
r.text = "3. COLLECTOR APP: "
r.font.bold = True
r.font.color.rgb = COLOR_WHITE
r.font.size = Pt(10)
r2 = p7_w.add_run()
r2.text = "Assigned pickup tasks, GIS coordinates & route navigation, before/after evidence photo upload, and one-click task completion."
r2.font.color.rgb = COLOR_ICE_BLUE
r2.font.size = Pt(9.5)

# Card 2: Operational Flowchart on Slide 7
card7_2 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, panel7_left, 4000000, panel7_width, 1950000)
card7_2.fill.solid()
card7_2.fill.fore_color.rgb = COLOR_BOX_BG
card7_2.line.color.rgb = COLOR_CYAN
card7_2.line.width = Pt(1)

tf7_2 = card7_2.text_frame
tf7_2.word_wrap = True
tf7_2.margin_left = Inches(0.12)
tf7_2.margin_right = Inches(0.12)
tf7_2.margin_top = Inches(0.08)
tf7_2.margin_bottom = Inches(0.08)

p7_2 = tf7_2.paragraphs[0]
p7_2.text = "VERIFIED INCIDENT LIFECYCLE"
p7_2.font.name = 'Courier New'
p7_2.font.size = Pt(10)
p7_2.font.bold = True
p7_2.font.color.rgb = COLOR_AMBER

flow_steps = [
    ("[1] Citizen Camera Scan", "HTML5 camera capture on smartphone"),
    ("[2] ResNet-18 Inference", "7-class prediction + confidence in ~22ms"),
    ("[3] Geotagged Report", "GPS lat/long + prediction saved in PostgreSQL"),
    ("[4] Admin Verification", "Municipal triage & dispatch prioritization"),
    ("[5] Collector Route", "GIS navigation to verified waste site"),
    ("[6] Evidence & Complete", "Post-cleanup photo upload & job closure")
]

for step_title, step_desc in flow_steps:
    p_step = tf7_2.add_paragraph()
    r_st = p_step.add_run()
    r_st.text = f"{step_title}: "
    r_st.font.bold = True
    r_st.font.color.rgb = COLOR_HIGHLIGHT
    r_st.font.size = Pt(9)
    r_sd = p_step.add_run()
    r_sd.text = step_desc
    r_sd.font.color.rgb = COLOR_ICE_BLUE
    r_sd.font.size = Pt(8.5)

# Save output presentation
prs.save(OUT_PPT)
print(f"Presentation regenerated successfully: {OUT_PPT}")
